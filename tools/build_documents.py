from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]

def base():
 d=Document(); s=d.sections[0]
 s.top_margin=Cm(2); s.bottom_margin=Cm(2); s.left_margin=Cm(2.5); s.right_margin=Cm(2)
 s.page_width=Cm(21); s.page_height=Cm(29.7)
 normal=d.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
 normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
 for name,size in [('Title',22),('Heading 1',16),('Heading 2',13)]:
  st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
 for border in list(d.styles.element.xpath('.//w:pBdr')):
  border.getparent().remove(border)
 return d

def p(d,text): return d.add_paragraph(text)
def h(d,text): return d.add_heading(text,1)
def code(d,text):
 para=d.add_paragraph(); para.paragraph_format.line_spacing=1.0
 para.paragraph_format.space_after=Pt(8)
 r=para.add_run(text); r.font.name='Consolas'; r.font.size=Pt(9)

def shot(d,name,caption,width=6.7):
 path=ROOT/'screenshots/lab1'/name
 if not path.exists(): raise FileNotFoundError(path)
 para=d.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.add_run().add_picture(str(path),width=Cm(width))
 para=d.add_paragraph(caption); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.paragraph_format.space_after=Pt(8)

def pair(d,a,b,captions):
 table=d.add_table(rows=1,cols=2)
 for cell,name,caption in zip(table.rows[0].cells,[a,b],captions):
  para=cell.paragraphs[0]; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
  para.add_run().add_picture(str(ROOT/'screenshots/lab1'/name),width=Cm(7.0))
  para=cell.add_paragraph(caption); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
  para.runs[0].font.size=Pt(10)

# Adapted assignment is a separate working document; the supplied original is unchanged.
d=base(); d.add_heading('Задание проекта на Flutter',0)
p(d,'Мобильный органайзер домашних растений. Вариант 7. Лабораторные работы 1–4. Заяц М.С., группа ИТИ-41.')
h(d,'Общие требования')
p(d,'Разработать мобильное приложение с декларативным интерфейсом Flutter на языке Dart. Основная платформа демонстрации — Android. Приложение запускается на реальном устройстве либо в эмуляторе Android Studio. Исходное задание допускает Flutter по согласованию с преподавателем.')
h(d,'Технические и архитектурные требования')
p(d,'Сохранить шаблон MVVM. Model описывает сущности и правила предметной области; ViewModel управляет состоянием, выполняет операции через слой данных и уведомляет View; View состоит из виджетов Flutter. Для работы 1 используются ChangeNotifier и ListenableBuilder. Зависимости передаются через конструкторы.')
p(d,'Соответствия технологий: Swift → Dart; SwiftUI → Flutter; SwiftData → SQLite; Realm → отдельное нереляционное хранилище; Codable → JSON-сериализация Dart; Combine → Streams и async/await; URLSession → HTTP-клиент; UserNotifications → локальные уведомления Android. Конкретные пакеты выбираются при выполнении соответствующего этапа.')
p(d,'Собственный REST-сервис не обязателен. Допускается тематический открытый API или mock-сервер. Для автоматизированного ввода названия растения требуется распознавание текста, а не штрихкода. При отсутствии камеры предусматривается mock с изображениями этикеток.')
h(d,'Требования к отчётам')
for t in ['Ссылка на открытый GitHub-репозиторий и снимки реализованных экранов и функций.',
 'Примеры кода интерфейса Flutter и реализации MVVM.',
 'В работе с хранением — примеры кода БД и сериализации; в сетевой работе — REST и асинхронные потоки.',
 'Примеры специальных архитектурных механизмов, если они предусмотрены этапом.',
 'Результаты проверок, ответы на контрольные вопросы в адаптации для Flutter и описание запуска.']:
 p(d,'• '+t)
d.add_page_break(); h(d,'Функциональная часть и этапы')
h(d,'Лабораторная работа 1')
p(d,'Зелёная часть. Цель — освоить декларативный интерфейс, компоновку, состояние и базовую навигацию во Flutter.')
for t in ['Отображение списка комнатных растений.','Интерактивный календарь процедур.','Визуальные индикаторы необходимости полива со сменой цвета карточки при просрочке.']: p(d,'• '+t)
p(d,'Данные на первом этапе демонстрационные, находятся в памяти. Просрочка определяется по текущей локальной дате. Требования к постоянному хранению и сети реализуются в следующих работах.')
h(d,'Лабораторная работа 2')
p(d,'Лиловая часть. Журнал ухода за растением: история поливов, подкормок и пересадок. Ботаническая справка по условиям содержания, свету и влажности. Сканирование текстового названия цветка с этикетки или ценника. Развитие разделения Model, View и ViewModel.')
h(d,'Лабораторная работа 3')
p(d,'Жёлтая часть. Основная реляционная база хранит карточки растений, журнал процедур и расписания триггеров. Вспомогательная нереляционная база хранит базовый справочник семейств растений и типов удобрений: уникальный идентификатор, название и иконка.')
h(d,'Лабораторная работа 4')
p(d,'Голубая часть. Распознанное название отправляется асинхронным REST-запросом в энциклопедию растений, например Perenual. Приложение получает базовый регламент ухода, сохраняет его в БД и генерирует цепочку локальных напоминаний. Обработка ошибок сети и отмена подписок выполняются средствами Dart.')
p(d,'Источник: «РПiPiP Задания на лабораторные работы.docx». Содержание варианта сохранено; технологические требования адаптированы для Flutter. Исходный файл не изменён.')
d.save(ROOT/'docs/assignment_flutter.docx')

if not (ROOT/'screenshots/lab1/01-catalog.png').exists():
 print('Assignment created; report awaits actual screenshots.'); raise SystemExit(0)
d=base(); d.add_heading('Отчёт по лабораторной работе 1',0)
p(d,'Проектирование и реализация пользовательского интерфейса средствами Flutter')
h(d,'Мобильный органайзер домашних растений')
p(d,'Вариант 7\nВыполнил: Заяц М.С.\nГруппа: ИТИ-41\n2026 год')
p(d,'Репозиторий: https://github.com/Jewilor/home-plants-organizer')
h(d,'Цель работы')
p(d,'Изучить декларативный подход к построению интерфейса, компоновку виджетов Flutter, управление состоянием и навигацию. Реализовать главный экран органайзера домашних растений.')
h(d,'Задание')
p(d,'Выполнить зелёную часть варианта 7: список комнатных растений; интерактивный календарь процедур; визуальные индикаторы необходимости полива со сменой цвета карточки при просрочке.')
h(d,'Применённые технологии')
p(d,'Dart 3.13.3, Flutter 3.47.4, Material 3, Android Studio 2026.1.4. Архитектура MVVM реализована моделями Plant и CareProcedure, классом GardenViewModel на ChangeNotifier и представлением GardenScreen. Данные на этом этапе находятся в памяти и задаются DemoGarden.')
p(d,'Полная адаптация общих и технических требований приведена в docs/assignment_flutter.docx. Журнал, OCR, базы данных, сеть и напоминания относятся к работам 2–4 и в работу 1 не включены.')
d.add_page_break(); h(d,'Каталог растений')
p(d,'Экран «Мой сад» содержит четыре демонстрационных растения. Для каждого указаны название, вид, комната и срок полива. Карточка монстеры с просроченным сроком имеет красный фон; фикус с поливом сегодня — янтарный; будущий срок обозначен зелёным. Цвет дублируется текстом и иконкой.')
pair(d,'01-catalog.png','02-catalog-scroll.png',['Рисунок 1 — Начало каталога','Рисунок 2 — Нижняя часть списка'])
p(d,'Срок сравнивается с текущей локальной датой без времени суток. Выбор другой даты в календаре не меняет статус полива. При смене суток и возвращении приложения на экран статус обновляется.')
d.add_page_break(); h(d,'Интерактивный календарь')
p(d,'Вкладка «Календарь» позволяет выбрать день и перейти между месяцами. Точка отмечает даты с процедурами. Кнопка «Сегодня» возвращает текущую дату. Для выбранного дня отображаются тип процедуры и растение; свободный день имеет отдельное пустое состояние.')
pair(d,'03-calendar.png','04-calendar-empty.png',['Рисунок 3 — Процедуры выбранного дня','Рисунок 4 — День без процедур'])
p(d,'Полив, подкормка и пересадка задаются демонстрационным расписанием. Для перехода между месяцами используются календарные значения DateTime, включая смену года и високосный февраль.')
d.add_page_break(); h(d,'Примеры реализации')
p(d,'Model. Правило срока полива не зависит от виджетов. Файл lib/models/plant.dart.')
code(d,"WateringStatus statusAt(DateTime now) {\n  final due = dateOnly(nextWatering);\n  final today = dateOnly(now);\n  if (due.isBefore(today)) return WateringStatus.overdue;\n  if (due == today) return WateringStatus.today;\n  return WateringStatus.upcoming;\n}")
p(d,'ViewModel. Выбор дня меняет состояние и уведомляет представление. Файл lib/viewmodels/garden_view_model.dart.')
code(d,"void selectDay(DateTime date) {\n  selectedDate = dateOnly(date);\n  visibleMonth = DateTime(date.year, date.month);\n  notifyListeners();\n}\n\nList<CareProcedure> proceduresOn(DateTime date) =>\n    procedures.where((p) => sameDay(p.date, date)).toList();")
p(d,'View. ListenableBuilder перестраивает интерфейс при изменениях ViewModel. Файл lib/views/garden_screen.dart; ниже сокращённый фрагмент привязки.')
code(d,"ListenableBuilder(\n  listenable: model,\n  builder: (context, _) => Scaffold(\n    // Каталог и календарь читают состояние model.\n  ),\n)")
p(d,'Иллюстрации растений рисуются CustomPainter. Русская локаль задаётся в MaterialApp через flutter_localizations. Постоянное хранилище и REST-клиент на данном этапе не реализованы; соответствующие примеры будут включены в следующие отчёты.')
h(d,'Проверки')
p(d,'flutter analyze: замечаний нет. flutter test: 7 тестов пройдено. Проверены сроки на границе года, високосный февраль, фильтрация процедур, смена суток, выбор дня, переход месяца, возврат к сегодня, цвет просроченной карточки, размеры 320 и 1024 пикселя при увеличении текста в 1,5 раза.')
p(d,'Android APK собран и запущен в эмуляторе Plants_API_36. Снимки получены из работающего приложения. Каталог прокручен, вкладка календаря открыта, выбран день без процедур.')
d.add_page_break(); h(d,'Контрольные вопросы в адаптации для Flutter')
questions=[
('1. Императивный и декларативный интерфейс','В императивном подходе разработчик изменяет отдельные элементы интерфейса. Во Flutter метод build описывает дерево виджетов для текущего состояния, а фреймворк обновляет соответствующие элементы и объекты отрисовки. Это соответствует идее сравнения UIKit и SwiftUI в исходном вопросе.'),
('2. Жизненный цикл представления','Widget — неизменяемая конфигурация. Для StatefulWidget создаётся State: initState выполняется один раз, build может вызываться многократно, didUpdateWidget обрабатывает новую конфигурацию, dispose освобождает ресурсы. В проекте таймер и наблюдение за жизненным циклом отключаются в dispose.'),
('3. Контейнеры компоновки','Row размещает дочерние виджеты горизонтально, Column — вертикально, Stack — слоями. Это аналоги HStack, VStack и ZStack. Expanded распределяет свободное пространство внутри Row или Column.'),
('4. Порядок оформления','Оформление выражается вложенными виджетами. Padding снаружи Container добавляет внешний отступ, а внутри — отступ содержимого. Поэтому перестановка обёрток меняет ограничения, фон и область взаимодействия.'),
('5. Навигация','Для стека экранов применяются Navigator и маршруты, например MaterialPageRoute. В работе 1 NavigationBar переключает два представления одного главного экрана; состояние календаря хранится в общей ViewModel. Детальные маршруты потребуются во второй работе.'),
('6. Локальное состояние и привязка','State и setState используются для локального состояния виджета. Передача значения и callback родителем позволяет дочернему виджету запросить изменение, подобно назначению Binding. Номер вкладки хранится в State, выбранный день — в ViewModel.'),
('7. Наблюдение за внешними изменениями','ChangeNotifier вызывает notifyListeners после изменения состояния. ListenableBuilder подписывается на него и перестраивает представление. Это используемый в проекте механизм реактивного обновления вместо Observable из SwiftUI.'),
('8. Безопасная область','SafeArea учитывает системные панели, вырезы и другие отступы устройства, чтобы содержимое не перекрывалось ими. В проекте основной интерфейс помещён в SafeArea.'),
('9. Адаптивный интерфейс','LayoutBuilder сообщает доступную ширину. На телефоне представления переключаются вкладками, от 850 пикселей каталог и календарь располагаются рядом. Прокрутка, Expanded и масштабирование чисел календаря предотвращают переполнения.'),
('10. Передача общих зависимостей','Зависимости можно передавать через конструкторы, InheritedWidget или специализированный контейнер. Здесь GardenScreen принимает GardenViewModel, а ViewModel принимает clock и источник данных. Это позволяет тестам задавать фиксированную дату.')]
for title,answer in questions[:5]: d.add_heading(title,2); p(d,answer)
d.add_page_break(); h(d,'Контрольные вопросы и результат')
for title,answer in questions[5:]: d.add_heading(title,2); p(d,answer)
h(d,'Результат')
p(d,'Реализована зелёная часть варианта 7: русскоязычный каталог, календарь процедур и индикаторы полива. Получена основа для дальнейшего развития в лабораторных работах 2–4.')
d.save(ROOT/'reports/lab1_report.docx')
print('Created assignment and lab 1 report.')
