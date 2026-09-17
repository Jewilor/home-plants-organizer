from docx import Document
from pathlib import Path
from pypdf import PdfReader
import json,hashlib
path=Path('D:/3 курс/6 семестр/НейрАналИз/лаб 5/lab5.docx')
d=Document(path)
print('SECTIONS',len(d.sections),'TABLES',len(d.tables))
for s in d.sections:
 print('SECTION',[(k,str(getattr(s,k))) for k in ['page_width','page_height','left_margin','right_margin','top_margin','bottom_margin','header_distance','footer_distance']])
for i,p in enumerate(d.paragraphs):
 if i<65 or p.style.name.startswith('Heading'):
  print(i,repr(p.text),p.style.name,'align',p.alignment,'first',p.paragraph_format.first_line_indent,'before',p.paragraph_format.space_before,'after',p.paragraph_format.space_after,'line',p.paragraph_format.line_spacing)
  if p.text: print('RUNS',[(r.text[:70],r.font.name,r.font.size.pt if r.font.size else None,r.bold) for r in p.runs[:4]])
for i,page in enumerate(PdfReader('reports/qa/reference.pdf').pages):
 print('PAGE',i+1,(page.extract_text() or '')[:300].replace('\n',' '))
print('SHA256',hashlib.sha256(path.read_bytes()).hexdigest())
