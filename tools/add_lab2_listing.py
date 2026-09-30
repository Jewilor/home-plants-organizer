"""Append the complete lab 2 source listing without altering the report body."""
from __future__ import annotations

import hashlib
import json
import subprocess
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from build_lab2_report import blank, style_paragraph

ROOT = Path(__file__).resolve().parents[1]
BASE_REVISION = '7f6c856'
LAB2_REVISION = '7563070'
SOURCE = ROOT / 'reports/Лабораторная работа №2.docx'
OUTPUT = ROOT / 'reports/Лабораторная работа №2 с листингом.docx'
TEXT_SOURCE = SOURCE.with_suffix('.md')
TEXT_OUTPUT = OUTPUT.with_suffix('.md')
QA = ROOT / 'reports/qa/lab2-listing'
FILES = (
    'pubspec.yaml',
    'assets/botanical_profiles.json',
    'lib/models/plant.dart',
    'lib/models/care_record.dart',
    'lib/models/botanical_profile.dart',
    'lib/services/botanical_repository.dart',
    'lib/services/label_recognition_service.dart',
    'lib/viewmodels/garden_view_model.dart',
    'lib/viewmodels/plant_detail_view_model.dart',
    'lib/viewmodels/label_scan_view_model.dart',
    'lib/views/editors.dart',
    'lib/views/garden_screen.dart',
    'lib/views/plant_detail_screen.dart',
    'lib/views/label_scan_screen.dart',
    'lib/main.dart',
    'test/lab2_test.dart',
    'integration_test/lab2_demo_test.dart',
    'test_driver/lab2_screenshots.dart',
)
INTRO = (
    'В приложении приведены полные тексты файлов, созданных или изменённых '
    'при выполнении второй лабораторной работы. Включены также настройки '
    'зависимостей, данные справочника и программы автоматической проверки. '
    'В изменённых файлах сохранены необходимые части кода первой лабораторной '
    'работы, которые используются при работе приложения.'
)


def git(*args: str) -> str:
    return subprocess.check_output(
        ['git', *args], cwd=ROOT, encoding='utf-8'
    )


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def selected_files() -> set[str]:
    changed = git('diff', '--name-only', BASE_REVISION, LAB2_REVISION).splitlines()
    return {
        name for name in changed
        if (name.startswith(('lib/', 'test/', 'integration_test/', 'test_driver/'))
            and name.endswith('.dart'))
        or name in ('pubspec.yaml', 'assets/botanical_profiles.json')
    }


def verify_paragraph_style(paragraph, size: int, code: bool = False) -> None:
    f = paragraph.paragraph_format
    assert f.space_before.pt == 0 and f.space_after.pt == 0
    assert f.line_spacing == 1
    for run in paragraph.runs:
        assert run.font.name == 'Times New Roman'
        assert run.font.size.pt == size
        if code:
            assert run.italic is True
    mark = paragraph._p.find(qn('w:pPr')).find(qn('w:rPr'))
    assert mark.find(qn('w:sz')).get(qn('w:val')) == str(2 * size)


def build() -> None:
    assert selected_files() == set(FILES), 'Listing must cover every lab 2 source file.'
    originals = {SOURCE: sha256(SOURCE), ROOT / 'отчет 1.docx': sha256(ROOT / 'отчет 1.docx')}
    contents = {name: git('show', f'{LAB2_REVISION}:{name}') for name in FILES}
    appendix = Document()
    paragraph_checks = []

    for index, text in enumerate(('Приложение А', '(обязательное)', 'Листинг программы')):
        p = appendix.add_paragraph(text)
        style_paragraph(p, center=True, indent=False, bold=True)
        p.paragraph_format.page_break_before = index == 0
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        paragraph_checks.append((p, 14, False))
    p = blank(appendix, keep=True)
    paragraph_checks.append((p, 14, False))
    p = appendix.add_paragraph(INTRO)
    style_paragraph(p)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    paragraph_checks.append((p, 14, False))
    p = blank(appendix, keep=True)
    paragraph_checks.append((p, 14, False))

    offsets = {}
    for file_index, name in enumerate(FILES):
        if file_index:
            p = blank(appendix, keep=True)
            paragraph_checks.append((p, 14, False))
        p = appendix.add_paragraph(f'«{name}»')
        style_paragraph(p, bold=True)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        paragraph_checks.append((p, 14, False))
        p = blank(appendix, keep=True)
        paragraph_checks.append((p, 14, False))
        start = len(appendix.paragraphs)
        for line in contents[name].splitlines():
            # Preserve indentation, operators and source text exactly.
            p = appendix.add_paragraph(line)
            style_paragraph(p, size=11 if line else 14, code=bool(line))
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.widow_control = True
            paragraph_checks.append((p, 11 if line else 14, bool(line)))
        offsets[name] = (start, len(appendix.paragraphs))

    for p, size, code in paragraph_checks:
        verify_paragraph_style(p, size, code)
    for name, (start, end) in offsets.items():
        actual = [p.text for p in appendix.paragraphs[start:end]]
        assert actual == contents[name].splitlines(), f'Incomplete listing: {name}'

    with ZipFile(SOURCE) as original:
        parts = {entry.filename: original.read(entry.filename) for entry in original.infolist()}
        document = etree.fromstring(parts['word/document.xml'])
        body = document.find(qn('w:body'))
        old_elements = [etree.tostring(e) for e in body]
        old_paragraph_count = len(Document(SOURCE).paragraphs)
        for p in appendix.paragraphs:
            body.insert(len(body) - 1, deepcopy(p._p))
        document_bytes = etree.tostring(document, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(OUTPUT, 'w') as output:
            for entry in original.infolist():
                output.writestr(entry, document_bytes if entry.filename == 'word/document.xml' else parts[entry.filename])

    with ZipFile(OUTPUT) as output:
        assert set(output.namelist()) == set(parts)
        for name, content in parts.items():
            if name != 'word/document.xml':
                assert output.read(name) == content, f'Original package part changed: {name}'
        new_body = etree.fromstring(output.read('word/document.xml')).find(qn('w:body'))
        assert [etree.tostring(e) for e in list(new_body)[:len(old_elements)-1]] == old_elements[:-1]
        assert etree.tostring(new_body[-1]) == old_elements[-1]

    final_doc = Document(OUTPUT)
    final_appendix = final_doc.paragraphs[old_paragraph_count:]
    assert len(final_appendix) == len(appendix.paragraphs)
    for name, (start, end) in offsets.items():
        assert [p.text for p in final_appendix[start:end]] == contents[name].splitlines()
    assert all(sha256(path) == digest for path, digest in originals.items())

    text = TEXT_SOURCE.read_text(encoding='utf-8').rstrip()
    text += '\n\n**Приложение А**\n\n**(обязательное)**\n\n**Листинг программы**\n\n' + INTRO + '\n'
    for name in FILES:
        language = Path(name).suffix.lstrip('.')
        text += f'\n**«{name}»**\n\n```{language}\n{contents[name].rstrip()}\n```\n'
    TEXT_OUTPUT.write_text(text, encoding='utf-8')

    QA.mkdir(parents=True, exist_ok=True)
    manifest = {
        'base_revision': BASE_REVISION,
        'lab2_revision': LAB2_REVISION,
        'output': OUTPUT.name,
        'file_count': len(FILES),
        'line_count': sum(len(value.splitlines()) for value in contents.values()),
        'original_report_paragraphs': old_paragraph_count,
        'appendix_paragraphs': len(final_appendix),
        'original_report_sha256': originals[SOURCE],
        'reference_sha256': originals[ROOT / 'отчет 1.docx'],
        'files': {name: {'lines': len(contents[name].splitlines()),
                         'sha256': hashlib.sha256(contents[name].encode('utf-8')).hexdigest()}
                  for name in FILES},
        'checks': 'Exact source text; full changed source coverage; original body and all other package parts preserved; paragraph styles verified.',
    }
    (QA / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(OUTPUT), 'files': len(FILES), 'lines': manifest['line_count'],
                      'paragraphs': len(final_doc.paragraphs)}, ensure_ascii=False))


if __name__ == '__main__':
    build()
