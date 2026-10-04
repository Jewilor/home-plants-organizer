"""Build the separate lab 4 report from the retained university report design."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json
import re
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from build_lab2_report import style_paragraph, blank

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'reports/Лабораторная работа №4.docx'
MARKDOWN = REPORT.with_suffix('.md')
REFERENCE = ROOT / 'reports/Лабораторная работа №3 с листингом А.docx'
FENCE = chr(96) * 3
LISTING_FILES = ['pubspec.yaml', 'lib/models/care_guide.dart', 'lib/models/reminder.dart', 'lib/models/plant.dart', 'lib/models/garden_snapshot.dart', 'lib/services/plant_encyclopedia.dart', 'lib/services/reminder_service.dart', 'lib/services/android_reminder_service.dart', 'lib/services/sqlite_garden_repository.dart', 'lib/services/garden_resources.dart', 'lib/services/storage_bootstrap_native.dart', 'lib/services/storage_bootstrap_web.dart', 'lib/viewmodels/care_guide_view_model.dart', 'lib/viewmodels/garden_view_model.dart', 'lib/views/care_guide_screen.dart', 'lib/views/reminders_screen.dart', 'lib/views/plant_detail_screen.dart', 'lib/views/label_scan_screen.dart', 'lib/views/garden_screen.dart', 'lib/views/editors.dart', 'android/app/src/main/AndroidManifest.xml', 'android/app/build.gradle.kts', 'android/app/src/main/res/drawable/ic_notification.xml', 'demo_api/encyclopedia.json', 'test/lab4_test.dart', 'test/lab3_storage_test.dart', 'integration_test/lab4_demo_test.dart', 'test_driver/lab4_screenshots.dart']
APPENDIX_TITLE = 'Приложение А'
GITHUB_URL = 'https://github.com/Jewilor/home-plants-organizer'


def listing_sources():
    # Full final source of every app/configuration/test file authored or modified for lab 4.
    # Generated lockfile, report builder and the user's independent demo comment are excluded.
    return {path: (ROOT / path).read_text(encoding='utf-8-sig').rstrip()
            for path in LISTING_FILES}


def extract(relative, begin=None, end=None):
    source = (ROOT / relative).read_text(encoding='utf-8')
    a = source.index(begin) if begin else 0
    b = source.index(end, a) if end else len(source)
    return source[a:b].rstrip()


def report_text():
    content = (ROOT / 'docs/lab4_report_text.md').read_text(encoding='utf-8')
    sources = listing_sources()
    content += '\n\nПриложение А\n\n(обязательное)\n\nЛистинг программы\n'
    for index, (path, source) in enumerate(sources.items(), 1):
        content += f'\nЛистинг А.{index} – {path}\n\n{FENCE}\n{source}\n{FENCE}\n'
    return content


def hyperlink(p, label, url):
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    for tag, attributes in [
        ('w:rFonts', {'w:ascii':'Times New Roman','w:hAnsi':'Times New Roman','w:cs':'Times New Roman'}),
        ('w:sz', {'w:val':'28'}),
        ('w:i', {}),
        ('w:color', {'w:val':'000000'}),
        ('w:u', {'w:val':'single'}),
    ]:
        element = OxmlElement(tag)
        for key, value in attributes.items():
            element.set(qn(key), value)
        props.append(element)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    node.append(run)
    p._p.append(node)


def add_text(document, line):
    text = line.replace('**', '')
    if line.startswith('**') or line.startswith('Вывод:'):
        if document.paragraphs[-1].text.strip():
            blank(document)
    p = document.add_paragraph()
    if line.startswith('**') and ':' in text:
        label, rest = text.split(':', 1)
        p.add_run(label + ':').bold = True
        p.add_run(rest)
    else:
        cursor = 0
        for match in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)', text):
            p.add_run(text[cursor:match.start()])
            hyperlink(p, match[1], match[2])
            cursor = match.end()
        p.add_run(text[cursor:])
    style_paragraph(p)
    p.paragraph_format.widow_control = True
    return p


def build():
    before = sha256(REFERENCE.read_bytes()).hexdigest()
    content = report_text()
    MARKDOWN.write_text(content, encoding='utf-8')
    d = Document(REFERENCE)
    title = [deepcopy(p._p) for p in d.paragraphs[:36]]
    body = d._element.body
    for item in list(body):
        if not item.tag.endswith('sectPr'):
            body.remove(item)
    for element in title:
        body.insert(len(body)-1, element)
    d.paragraphs[15].text = 'ЛАБОРАТОРНАЯ РАБОТА № 4'
    d.paragraphs[18].text = 'на тему: «Сетевое взаимодействие на основе REST и реактивное обновление интерфейса средствами Flutter»'
    for i, p in enumerate(d.paragraphs):
        style_paragraph(p, indent=False)
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.keep_together = False
        if i in (25,26,27,28):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(8)
        elif p.text:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section = d.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.top_margin = section.bottom_margin = Cm(2)
    for style in d.styles:
        if style.type == 1:
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing = 1
    d.styles['Normal'].font.name = 'Times New Roman'
    d.styles['Normal'].font.size = Pt(14)
    in_code = False
    started = False
    figures = 0
    appendix = False
    listing_count = 0
    for line in content.splitlines():
        if not started:
            if line.startswith('**Цель работы:'):
                started = True
            else:
                continue
        if line.startswith(FENCE):
            in_code = not in_code
            continue
        if in_code:
            p = d.add_paragraph(line)
            style_paragraph(p, indent=False, size=11 if line.strip() else 14, code=bool(line.strip()))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_together = True
            p.paragraph_format.widow_control = False
            continue
        if not line.strip():
            continue
        if line == APPENDIX_TITLE:
            appendix = True
            p = d.add_paragraph(line)
            style_paragraph(p, center=True, indent=False)
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            continue
        if appendix and line in ('(обязательное)', 'Листинг программы'):
            p = d.add_paragraph(line)
            style_paragraph(p, center=True, indent=False)
            p.paragraph_format.keep_with_next = True
            continue
        if appendix and line.startswith('Листинг А.'):
            if d.paragraphs[-1].text.strip():
                blank(d)
            listing_count += 1
            p = d.add_paragraph(line)
            style_paragraph(p, indent=False, size=11, code=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            continue
        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            blank(d, keep=True)
            p = d.add_paragraph()
            style_paragraph(p, center=True, indent=False)
            picture = p.add_run().add_picture(str((ROOT/'reports'/image[2]).resolve()), width=Cm(5.8))
            picture._inline.docPr.set('descr', image[1])
            p.paragraph_format.keep_with_next = True
            blank(d, keep=True)
            caption = d.add_paragraph(image[1])
            style_paragraph(caption, center=True, indent=False)
            caption.paragraph_format.keep_together = True
            blank(d)
            figures += 1
        else:
            p = add_text(d, line)
            if line.endswith('.dart.'):
                p.paragraph_format.keep_with_next = True
    d.core_properties.title = 'Лабораторная работа №4. Мобильный органайзер домашних растений'
    d.core_properties.subject = 'Асинхронное сетевое взаимодействие и локальные напоминания'
    d.core_properties.author = 'Заяц М. С.'
    d.core_properties.last_modified_by = 'Заяц М. С.'
    drawing_id = 1
    parts = [d.part]
    parts.extend(part for part in d.part.related_parts.values() if hasattr(part, 'element'))
    for part in parts:
        for node in part.element.iter():
            if node.tag == '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr':
                node.set('id', str(drawing_id))
                drawing_id += 1
    d.save(REPORT)
    assert sha256(REFERENCE.read_bytes()).hexdigest() == before
    assert figures == 10
    assert listing_count == len(LISTING_FILES)
    assert any(rel.reltype == RT.HYPERLINK and rel.target_ref == GITHUB_URL for rel in d.part.rels.values())
    print(json.dumps({'report':str(REPORT),'figures':figures,'paragraphs':len(d.paragraphs),'listings':listing_count,'reference_unchanged':True},ensure_ascii=False))


if __name__ == '__main__':
    build()
