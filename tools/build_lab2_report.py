from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def style_paragraph(p, *, center=False, indent=True, size=14, bold=False, code=False):
    f = p.paragraph_format
    f.space_before = Pt(0)
    f.space_after = Pt(0)
    f.line_spacing = 1
    f.first_line_indent = Cm(1.25 if indent else 0)
    f.left_indent = Cm(0)
    f.right_indent = Cm(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif indent:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0, 0, 0)
        if bold: r.bold = True
        if code: r.italic = True
    # Format the paragraph mark too, including empty spacing paragraphs.
    ppr = p._p.get_or_add_pPr()
    end_props = ppr.find(qn('w:rPr'))
    if end_props is None:
        end_props = OxmlElement('w:rPr')
        ppr.append(end_props)
    for tag, attrs in [('w:rFonts', {'w:ascii': 'Times New Roman', 'w:hAnsi': 'Times New Roman', 'w:cs': 'Times New Roman'}), ('w:sz', {'w:val': str(size * 2)}), ('w:szCs', {'w:val': str(size * 2)})]:
        element = end_props.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            end_props.append(element)
        for name, value in attrs.items():
            element.set(qn(name), value)
    # New Latin runs preserve weight, font and size of the original run.
    for run in list(p.runs):
        if code: continue
        text = run.text
        matches = list(re.finditer(r'[A-Za-z][A-Za-z0-9_/.:#?=&+%~-]*', text))
        if not matches: continue
        props = deepcopy(run._r.rPr)
        anchor = run._r
        segments = []
        last = 0
        for m in matches:
            if m.start() > last: segments.append((text[last:m.start()], False))
            segments.append((m.group(), True)); last = m.end()
        if last < len(text): segments.append((text[last:], False))
        for value, latin in segments:
            new = p.add_run(value)
            if props is not None: new._r.insert(0, deepcopy(props))
            if latin: new.italic = True
            anchor.addprevious(new._r)
        p._p.remove(anchor)


def blank(d, keep=False):
    p = d.add_paragraph()
    p.add_run('')
    style_paragraph(p, indent=False)
    p.paragraph_format.keep_with_next = keep
    return p


def build():
    d = Document(ROOT / 'docs/lab_report_template.docx')
    source = [deepcopy(p._p) for p in d.paragraphs[:44]]
    body = d._element.body
    for element in list(body):
        if element.tag.endswith('sectPr'): continue
        body.remove(element)
    indices = [0,1,2,3,4,5,6,8,9,10,11,12,13,14,15,17,18,19,20,
               21,22,23,24,25,26,33,34,35,36,37,38,39,40,41,43]
    for i in indices: body.insert(len(body)-1, deepcopy(source[i]))
    changes = {
        17: 'ЛАБОРАТОРНАЯ РАБОТА № 2',
        19: 'по дисциплине: «Разработка приложений для iPhone и iPad»',
        20: 'на тему: «Архитектурное проектирование мобильного приложения на основе шаблона Model–View–ViewModel»',
        33: 'Выполнил: студент группы ИТИ-41', 34: 'Заяц М. С.',
        35: 'Принял: старший преподаватель', 36: 'Семенченя Т. С.',
    }
    for j, i in enumerate(indices):
        p = d.paragraphs[j]
        if i in changes: p.text = changes[i]
        style_paragraph(p, indent=False)
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.keep_together = False
        if i in [33,34,35,36]:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(8)
        elif p.text: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sec = d.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3); sec.right_margin = Cm(1)
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    d.styles['Normal'].font.name = 'Times New Roman'
    d.styles['Normal'].font.size = Pt(14)
    for st in d.styles:
        if st.type == 1:
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(0)
            st.paragraph_format.line_spacing = 1
    # The reference title ends with its own page break.
    lines = (ROOT/'reports/Лабораторная работа №2.md').read_text(encoding='utf-8').splitlines()
    in_code = False
    code_lines = []
    started = False
    for line in lines:
        if not started:
            if line.startswith('**Цель работы:'): started=True
            else: continue
        if line.startswith('```'):
            if not in_code:
                in_code=True; code_lines=[]
            else:
                p=d.add_paragraph('\n'.join(code_lines))
                style_paragraph(p, indent=False, size=11, code=True)
                p.paragraph_format.keep_together=True
                in_code=False
            continue
        if in_code: code_lines.append(line); continue
        if not line.strip(): continue
        image=re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            path=(ROOT/'reports'/image.group(2)).resolve()
            if not path.exists(): raise FileNotFoundError(path)
            blank(d, keep=True)
            p=d.add_paragraph(); style_paragraph(p, center=True, indent=False)
            p.add_run().add_picture(str(path), width=Cm(6.4))
            p.paragraph_format.keep_with_next=True
            blank(d, keep=True)
            caption=d.add_paragraph(image.group(1))
            style_paragraph(caption, center=True, indent=False)
            caption.paragraph_format.keep_together=True
            blank(d)
        else:
            text=line.replace('**','')
            if (line.startswith('**') and not line.startswith('**Ход')) or line.startswith('Вывод:'):
                if d.paragraphs[-1].text.strip():
                    blank(d)
            p=d.add_paragraph()
            if line.startswith('**') and ':' in text:
                label, rest = text.split(':', 1)
                p.add_run(label + ':').bold = True
                p.add_run(rest)
            else:
                p.add_run(text)
            style_paragraph(p)
            if line.startswith('**'):
                label = text.split(':', 1)[0] + ':'
                offset = 0
                for run in p.runs:
                    if offset < len(label):
                        run.bold = True
                    offset += len(run.text)
    d.core_properties.title='Лабораторная работа 2 Мобильный органайзер домашних растений'
    d.core_properties.author='Заяц М. С.'
    d.core_properties.last_modified_by='Заяц М. С.'
    d.save(ROOT/'reports/Лабораторная работа №2.docx')


if __name__ == '__main__':
    build()
